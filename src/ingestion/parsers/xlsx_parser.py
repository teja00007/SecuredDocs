"""XLSX and CSV parsers."""

import unicodedata

from src.core.exceptions import IngestionError
from src.ingestion.parsers.base import BaseParser, ParseResult, Section


class XlsxParser(BaseParser):
    def supported_extensions(self) -> list[str]:
        return [".xlsx", ".xls"]

    def parse(self, file_path: str) -> ParseResult:
        try:
            import openpyxl
        except ImportError:
            raise IngestionError("openpyxl not installed")

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        except Exception as e:
            raise IngestionError(f"Failed to open XLSX: {e}")

        all_text: list[str] = []
        sections: list[Section] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            headers = [str(h) if h is not None else "" for h in rows[0]]
            header_line = "\t".join(headers)

            sheet_lines = [f"Sheet: {sheet_name}", header_line]
            for row in rows[1:]:
                cells = [str(c) if c is not None else "" for c in row]
                sheet_lines.append("\t".join(cells))

            sheet_text = "\n".join(sheet_lines)
            sheet_text = unicodedata.normalize("NFC", sheet_text)
            all_text.append(sheet_text)
            sections.append(Section(
                heading=sheet_name,
                body=sheet_text,
                level=1,
            ))

        wb.close()
        full_text = "\n\n".join(all_text)
        metadata = {"file_type": ".xlsx", "sheet_count": len(wb.sheetnames)}

        return ParseResult(text=full_text, metadata=metadata, sections=sections)


class CsvParser(BaseParser):
    def supported_extensions(self) -> list[str]:
        return [".csv"]

    def parse(self, file_path: str) -> ParseResult:
        import csv

        try:
            with open(file_path, "r", encoding="utf-8", errors="replace", newline="") as f:
                reader = csv.reader(f)
                rows = list(reader)
        except OSError as e:
            raise IngestionError(f"Failed to read CSV: {e}")

        if not rows:
            return ParseResult(text="", metadata={"file_type": ".csv"})

        lines = ["\t".join(rows[0])]  # header
        for row in rows[1:]:
            lines.append("\t".join(row))

        text = unicodedata.normalize("NFC", "\n".join(lines))
        metadata = {"file_type": ".csv", "row_count": len(rows) - 1}

        return ParseResult(text=text, metadata=metadata)
